from io import BytesIO
from pathlib import Path
import os
import pandas as pd
import pytest
from openpyxl import Workbook
from core.reader import read_report, SCHEMAS
from core.validation import validate,reconcile
from core.normalization import number,parse_date,identifier,safe_csv,DEFAULT_FLAGS
from core.metrics import metrics,ratio,previous_period,authorized_join,total

def fixture(kind,rows):
    w=Workbook(); s=w.active
    s.append(SCHEMAS[kind])
    s.append(['Entidad']); s.append(['Sucursal'])
    for row in rows:
        s.append(row)
    s.append(['SubTotal Sucursal']);s.append(['SubTotal Comercio'])
    out=BytesIO();w.save(out)
    return read_report(out.getvalue(),'prueba.xlsx',kind)

def pay(amount=10,ann='',loan='00001',day='01/10/2026'):
    return ['001','Persona','002',loan,1,'01',day,5,4,0,0,amount,'Efectivo','',ann]

def test_normalization():
    assert number('1.234,56')==1234.56
    assert number(1234.56)==1234.56
    assert pd.isna(number(''))
    assert pd.isna(parse_date('31/02/2026'))
    w=Workbook();c=w.active.cell(1,1,123);c.number_format='000000'
    assert identifier(c)=='000123'
    assert ratio(1,0) is None
    assert ratio(None,2) is None
    assert b"'=SUM" in safe_csv(pd.DataFrame({'a':['=SUM(1,2)']}))

def test_partial_duplicate_annulled_and_ambiguous():
    r=fixture('cobranzas',[pay(),pay(),pay(5,'Si'),pay(2,'desconocido')])
    d,q=validate(r)
    assert d.duplicado_potencial.sum()==2
    assert d.incluido.sum()==2
    assert metrics(d[d.incluido],'cobranzas')['Total cobrado']==20
    assert len(q)>0
    assert d.importe_no_explicado.iloc[0]==1
    strict=DEFAULT_FLAGS|{'':None}
    assert validate(r,strict)[0].incluido.sum()==0

def test_invalid_date_and_empty():
    r=fixture('cobranzas',[pay(day='31/02/2026')])
    d,_=validate(r)
    assert not d.incluido.any()
    assert total(d.iloc[:0],'cobrado') is None

def test_missing_amount_and_negative_are_visible():
    row=pay();row[7]=None
    negative=pay(-2,loan='00002')
    r=fixture('cobranzas',[row,negative])
    d,q=validate(r)
    assert total(d,'capital') is None
    assert total(d,'cobrado')==8
    assert q.motivo.str.contains('ausente|negativo').any()
    assert d.capital_original.iloc[0] is None or pd.isna(d.capital_original.iloc[0])

def test_duplicate_loans_and_cancelled():
    row=['001','Persona','002','00001','01/10/2026',3,'01',10,2,0,12,0,12,'Cancelado','Operador']
    r=fixture('prestamos',[row])
    assert validate(r)[0].incluido.all()
    r=fixture('prestamos',[row,row])
    assert not validate(r)[0].incluido.any()

def test_missing_columns():
    w=Workbook();w.active.append(['Cliente','Prestamo','Plan']);o=BytesIO();w.save(o)
    with pytest.raises(ValueError,match='faltan columnas'):
        read_report(o.getvalue(),'incompleto.xlsx','cobranzas')

def test_comparison():
    x=previous_period('2026-10-01','2026-10-19',pd.Timestamp('2026-01-01'),pd.Timestamp('2026-10-19'))
    assert x[1]==pd.Timestamp('2026-09-19')
    assert previous_period('2026-10-01','2026-10-19',pd.Timestamp('2026-10-01'),pd.Timestamp('2026-10-19')) is None
    assert previous_period('2026-10-02','2026-10-19',pd.Timestamp('2026-01-01'),pd.Timestamp('2026-10-19')) is None

def test_join_no_multiplication():
    loans=pd.DataFrame({'grupo_1':['A'],'prestamo':['001'],'capital':[100]})
    payments=pd.DataFrame({'grupo_1':['B','B','B'],'prestamo':['x','x','z'],'cobrado':[3,4,9]})
    mapping=pd.DataFrame({'entidad_origen':['B'],'prestamo_origen':['x'],'entidad_destino':['A'],'prestamo_destino':['001']})
    joined,unmatched=authorized_join(loans,payments,mapping)
    assert len(joined)==1 and joined.capital.sum()==100 and joined.cobrado.sum()==7
    assert len(unmatched)==1
    with pytest.raises(ValueError):
        authorized_join(loans,payments,pd.concat([mapping,mapping]))

FILES={
 'prestamos':'Prestamos Detalle_d6b361db-adcf-4e7e-a375-9219a6a54277.xlsx',
 'cobranzas':'Cobranzas Detalle_d4313e5e-c0b0-4d68-8208-5f3b503f97bb.xlsx',
 'cartera':'Saldo de la Cartera_9eaa877f-3f4f-4f93-aaac-9138cdb44dd7.xlsx'}

@pytest.mark.parametrize('kind,n,header,field,expected',[
 ('prestamos',15,11,'capital',1269999.96),
 ('cobranzas',10,13,'cobrado',256991.22),
 ('cartera',123,11,'saldo_capital',53482000)])
def test_attached_reports(kind,n,header,field,expected):
    path=Path(os.environ.get('REPORTS_DIR',str(Path.home()/'Downloads')))/FILES[kind]
    if not path.exists():
        pytest.skip('Archivo real no disponible en REPORTS_DIR')
    r=read_report(path.read_bytes(),path.name,kind)
    assert list(r.headers.values())==[header]
    assert len(r.data)==n
    assert r.data[field].sum()==pytest.approx(expected)
    c=reconcile(r)
    assert not c.estado.eq('Diferencia').any(),c[c.estado.eq('Diferencia')].to_dict('records')
    d,q=validate(r)
    assert d.prestamo.map(type).eq(str).all()
    assert d.entidad.notna().all() and d.sucursal.notna().all()
    if kind=='cartera':
        assert all(d['diferencia_'+v].abs().le(.02).all() for v in ['capital','interes','interes_iva','monto'])
    if kind=='cobranzas':
        assert d.incluido.sum()==10
        assert metrics(d[d.incluido],kind)['Importe no explicado']==pytest.approx(23416.56)
